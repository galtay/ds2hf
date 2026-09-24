from __future__ import annotations

import re
from pathlib import Path

import pytest
from ds2hf_pipeline.tcga import cdr


def _workbook(path: Path, sheets: dict[str, list[list]]) -> Path:
    """Write an .xlsx whose cells are exactly `sheets`; openpyxl stores
    "#N/A" and "#REF!" as error cells, as Excel does."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


_HEADER = [None, "bcr_patient_barcode", "OS.time", "tumor_status", "Redaction"]


def test_read_sheet_keeps_values_as_published(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "w.xlsx",
        {
            "S": [
                _HEADER,
                ["1", "TCGA-AA-0001", 100, "[Not Available]", "[Not Applicable]"],
                ["2", "TCGA-AA-0002", "#N/A", "#N/A", "Redacted"],
            ]
        },
    )
    table = cdr.read_sheet(path, "S")

    # The row-number column is dropped; every other header is verbatim, dots included.
    assert table.column_names == _HEADER[1:]
    assert str(table.schema.field("OS.time").type) == "int32"
    # `#N/A` error cells are null in integer and text columns alike...
    assert table.column("OS.time").to_pylist() == [100, None]
    assert table.column("tumor_status").to_pylist() == ["[Not Available]", None]
    # ...while placeholders stay distinct values. (The real workbook's empty
    # `Redaction` strings can't be written here: openpyxl stores "" as a
    # blank cell. `verify-cdr` covers them against the workbook itself.)
    assert table.column("Redaction").to_pylist() == ["[Not Applicable]", "Redacted"]


def test_read_sheet_refuses_other_excel_errors(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "w.xlsx", {"S": [_HEADER, ["1", "TCGA-AA-0001", "#REF!", "x", ""]]})
    with pytest.raises(ValueError, match="#REF!"):
        cdr.read_sheet(path, "S")


def test_read_sheet_refuses_mixed_columns(tmp_path: Path) -> None:
    """A fractional day count fits neither the integer nor the text rule."""
    path = _workbook(
        tmp_path / "w.xlsx",
        {"S": [_HEADER, ["1", "A", 1, "x", ""], ["2", "B", 1.5, "x", ""]]},
    )
    with pytest.raises(ValueError, match="OS.time"):
        cdr.read_sheet(path, "S")


def test_read_sheet_requires_the_row_number_column(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "w.xlsx", {"S": [_HEADER, ["7", "A", 1, "x", ""]]})
    with pytest.raises(ValueError, match="row number"):
        cdr.read_sheet(path, "S")


def test_read_notes_keeps_nesting(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "w.xlsx",
        {"N": [["Heading"], [None, "entry"], [None, None, "continued"], [None, None]]},
    )
    assert cdr.read_notes(path, "N") == [(0, "Heading"), (1, "entry"), (2, "continued")]


def test_cdr_card_names_no_other_hub_dataset(tmp_path: Path) -> None:
    """The card stands alone: a sibling dataset changing must never make it stale."""
    import pyarrow as pa
    import pyarrow.parquet as pq
    from ds2hf_pipeline.tcga.dataset_card import CDR_REPO_ID, write_cdr_card

    for config in cdr.CDR_CONFIGS:
        (tmp_path / config).mkdir()
        pq.write_table(pa.table({"type": ["BRCA"]}), tmp_path / config / "data.parquet")
    card = write_cdr_card(
        tmp_path,
        {config: 1 for config in cdr.CDR_CONFIGS},
        {config: [(0, "note")] for config in cdr.CDR_CONFIGS},
    ).read_text()

    assert "{" not in card.split("```")[0], "unrendered placeholder in the card"
    repos = set(re.findall(r"gabrielaltay/[\w.-]+", card))
    assert repos == {CDR_REPO_ID}
    assert "huggingface.co/datasets" not in card
